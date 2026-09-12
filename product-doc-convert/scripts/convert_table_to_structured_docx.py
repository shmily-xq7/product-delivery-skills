#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
表格转结构化DOCX文档
将表格内容转换为标题+正文的DOCX文档格式
"""

import argparse
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn


class TableToStructuredDocxConverter:
    """表格转结构化DOCX转换器"""
    
    def __init__(self, input_docx_path: str):
        """
        初始化转换器
        
        Args:
            input_docx_path: 输入DOCX文件路径
        """
        self.input_path = input_docx_path
        self.input_doc = Document(input_docx_path)
        self.output_doc = Document()
        
    def set_chinese_font(self, run, font_name='微软雅黑'):
        """
        设置中文字体
        
        Args:
            run: 文本运行对象
            font_name: 字体名称
        """
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    
    def add_title(self, text: str):
        """
        添加文档标题
        
        Args:
            text: 标题文本
        """
        title = self.output_doc.add_heading(text, level=1)
        title_run = title.runs[0]
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0, 0, 0)
        self.set_chinese_font(title_run)
        
    def add_heading_2(self, text: str):
        """
        添加二级标题
        
        Args:
            text: 标题文本
        """
        heading = self.output_doc.add_heading(text, level=2)
        heading_run = heading.runs[0]
        heading_run.font.size = Pt(14)
        heading_run.font.bold = True
        heading_run.font.color.rgb = RGBColor(31, 73, 125)
        self.set_chinese_font(heading_run)
        
    def add_heading_3(self, text: str):
        """
        添加三级标题
        
        Args:
            text: 标题文本
        """
        heading = self.output_doc.add_heading(text, level=3)
        heading_run = heading.runs[0]
        heading_run.font.size = Pt(12)
        heading_run.font.bold = True
        heading_run.font.color.rgb = RGBColor(79, 129, 189)
        self.set_chinese_font(heading_run)
        
    def add_paragraph(self, text: str):
        """
        添加正文段落
        
        Args:
            text: 段落文本
        """
        para = self.output_doc.add_paragraph(text)
        para_run = para.runs[0]
        para_run.font.size = Pt(11)
        self.set_chinese_font(para_run)
        
    def process_table_row(self, row_cells: list, current_h1: str):
        """
        处理表格行，转换为标题+正文格式
        
        Args:
            row_cells: 当前行的单元格列表
            current_h1: 当前的一级标题
            
        Returns:
            新的一级标题（如果有变化）
        """
        # 表格结构：第一列=一级标题 | 第二列=二级标题 | 第三列=正文
        if len(row_cells) >= 3:
            h1_text = row_cells[0].strip()
            h2_text = row_cells[1].strip()
            content = row_cells[2].strip()
            
            # 如果一级标题发生变化且不为空，添加新的一级标题
            if h1_text and h1_text != current_h1:
                self.add_heading_2(h1_text)
                current_h1 = h1_text
            
            # 二级标题（技术指标）
            if h2_text:
                self.add_heading_3(h2_text)
            
            # 正文内容
            if content:
                # 处理换行符
                content = content.replace('<br>', '\n')
                self.add_paragraph(content)
                
        return current_h1
    
    def convert(self) -> Document:
        """
        执行转换
        
        Returns:
            转换后的Document对象
        """
        # 添加文档主标题（从第一个段落获取）
        for para in self.input_doc.paragraphs:
            if para.text.strip():
                self.add_title(para.text.strip())
                break
        
        # 处理所有表格
        for table_idx, table in enumerate(self.input_doc.tables):
            if not table.rows:
                continue
            
            # 跟踪当前的一级标题
            current_h1 = ""
            
            # 处理数据行（跳过表头）
            for row_idx, row in enumerate(table.rows[1:], start=1):
                row_cells = [cell.text.strip() for cell in row.cells]
                current_h1 = self.process_table_row(row_cells, current_h1)
        
        return self.output_doc
    
    def save_to_file(self, output_path: str):
        """
        保存转换结果到文件
        
        Args:
            output_path: 输出文件路径
        """
        self.convert()
        
        # 确保输出目录存在
        output_dir = os.path.dirname(output_path)
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir)
        
        # 保存文档
        self.output_doc.save(output_path)
        
        print(f"✅ 转换完成！")
        print(f"📄 输入文件: {self.input_path}")
        print(f"📝 输出文件: {output_path}")
        print(f"📊 文档统计:")
        print(f"   - 原文档段落数: {len(self.input_doc.paragraphs)}")
        print(f"   - 原文档表格数: {len(self.input_doc.tables)}")
        print(f"   - 新文档段落数: {len(self.output_doc.paragraphs)}")


def main():
    """主函数"""
    parser = argparse.ArgumentParser(
        description="将 DOCX 中的表格转换为「标题 + 正文」的结构化 DOCX 文档。"
    )
    parser.add_argument("input_file", help="输入 DOCX 文件路径（其内容为表格）")
    parser.add_argument(
        "-o", "--output", dest="output_file", default=None,
        help="输出 DOCX 路径（默认在原文件旁生成 *_结构化.docx）",
    )
    args = parser.parse_args()

    input_file = args.input_file
    output_file = args.output_file or input_file.replace('.docx', '_结构化.docx')

    # 检查输入文件是否存在
    if not os.path.exists(input_file):
        print(f"❌ 错误: 文件不存在 - {input_file}")
        return
    
    try:
        # 创建转换器并执行转换
        converter = TableToStructuredDocxConverter(input_file)
        converter.save_to_file(output_file)
        
    except Exception as e:
        print(f"❌ 转换失败: {str(e)}")
        import traceback
        traceback.print_exc()


if __name__ == "__main__":
    main()
